"""Build the updated report from preserved main text and computed extension tables."""
from pathlib import Path
import csv, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parent
A = ROOT/'additional_robustness_checks'

def read(name):
    with (A/'tables'/name).open() as f:
        return list(csv.DictReader(f))
def f(x, digits=3):return f'{float(x):.{digits}f}'
def ptext(x):return 'p < .001' if float(x)<.001 else f'p = {float(x):.4f}'
def pcell(x):return '<.001' if float(x)<.001 else f'{float(x):.4f}'
def interval(r,lo='ci_low',hi='ci_high'):return f'[{f(r[lo])}, {f(r[hi])}]'

def extension_text():
    desc=read('anova_descriptives.csv');tests=read('anova_tests.csv');pairs=read('anova_posthoc.csv');loso=read('loso_estimates.csv')
    trust,accuracy=tests
    def pair(outcome,contrast):return next(r for r in pairs if r['outcome']==outcome and r['contrast']==contrast)
    tw=pair('Trust','Wide range - Point');aw=pair('Perceived accuracy','Wide range - Point')
    full=loso[0];omissions=loso[1:];largest=max(omissions,key=lambda r:abs(float(r['change_from_full'])))
    lowers=min(float(r['b']) for r in omissions);uppers=max(float(r['b']) for r in omissions)
    assert all(float(r['ci_low'])>0 for r in omissions)
    return f'''## 3 Additional analyses

The following student analyses extend the replication using the original data and exclusions. They are exploratory and are distinguished from the authors’ reported findings.

### Trust and perceived accuracy in Study 2

The first analysis asks whether estimate formats influence trust and perceived accuracy in similar ways. Unlike the composite outcome, separate items can reveal differences between these evaluations. After the original duplicate exclusions, the sample contains {trust['n']} participants. Point, moderate-range, and wide-range conditions contain {desc[0]['n']}, {desc[1]['n']}, and {desc[2]['n']} participants, respectively; neither item has missing responses. Table B1 reports group means and standard deviations, and Figure B1 shows individual 95% t confidence intervals.

Separate classical one-way ANOVAs detected condition differences in trust, F({int(float(trust['df_between']))}, {int(float(trust['df_within']))}) = {f(trust['F'])}, {ptext(trust['p'])}, and perceived accuracy, F({int(float(accuracy['df_between']))}, {int(float(accuracy['df_within']))}) = {f(accuracy['F'])}, {ptext(accuracy['p'])}. Both remain significant after Holm correction across the two outcomes. Because the median-centered Levene test indicated unequal variances for accuracy ({ptext(accuracy['levene_p'])}), Welch ANOVAs were also estimated; both confirmed the omnibus conclusions (Table B2). Pairwise comparisons use Welch t tests with Holm correction over all six comparisons; Table B3 also reports Bonferroni simultaneous 95% intervals.

Moderate ranges produced higher trust and accuracy ratings than both alternatives. Wide ranges exceeded point estimates on accuracy (difference = {f(aw['difference'])}; Holm {ptext(aw['p_holm_six_tests'])}), but the corresponding trust difference was inconclusive (difference = {f(tw['difference'])}; Holm {ptext(tw['p_holm_six_tests'])}). Thus, the observed patterns broadly favor moderate ranges, while wider estimates’ accuracy advantage does not translate into detectable higher trust. This describes the two outcomes’ patterns; it does not formally test equality of their treatment effects or show that the nonsignificant trust effect is zero.

This analysis overlaps with the authors’ Study 2 exploratory item-level OLS regressions (Study 2.do, section 3.4). ANOVA uses the same underlying group-mean model; the explicit omnibus tests, multiplicity adjustments, Welch sensitivity checks, and figure supplement that existing analysis. It is not a new substantive hypothesis. Supplement 2’s percentage/error accuracy measures differ from the sb_accurate item used here.

### Pooled study fixed effects and leave one study out

The second analysis tests whether the pooled range advantage depends on any single experiment. Comparable point and moderate-range conditions from Studies 1, 2, 3, 5, 6, and 7 were pooled, retaining the original five-item satisfaction composite and exclusion rules. Study 7 contributes its accurate point and accurate range conditions. Study 3 varies estimates across trials. The full sample contains {int(full['n']):,} participants, with no missing composites. OLS regresses satisfaction on a range indicator and study fixed effects, using conventional OLS standard errors and residual-degrees-of-freedom t intervals. These standard errors assume equal error variance across observations. The identical specification is then estimated six times, omitting one study each time (Table B4 and Figure B2).

The full-sample range coefficient is {f(full['b'])} points (SE = {f(full['se_ols'])}; 95% CI {interval(full)}; {ptext(full['p'])}). Omission estimates range from {lowers:.3f} to {uppers:.3f}; every confidence interval remains above zero. The largest absolute change occurs when Study {int(float(largest['omitted_study']))} is excluded (change = {float(largest['change_from_full']):+.3f}; coefficient = {f(largest['b'])}). The evidence therefore does not indicate that any single included study drives the positive pooled effect.

These related sensitivity models are not independent confirmations and do not establish identical effects across studies. Pooling also assumes independent observations across experiments; cross-study participant overlap cannot be checked using study-specific de-identified IDs. The check addresses study influence rather than field generalizability or every form of heterogeneity. No equivalent six-study fixed-effects leave-one-study-out analysis was found in the reviewed paper, complete supplement, or 11 supplied do-files. The supplement’s app-domain fixed effects concern a different design. This absence claim is limited to the materials reviewed.
'''

ext=extension_text()
(A/'additional_analyses.md').write_text(ext+'\n')
source=(ROOT/'report_source.md').read_text().replace('<!-- ADDITIONAL_ANALYSES -->',ext)
# Markdown keeps original main-study tables and adds live extension tables below.
def mdtable(heads,rows):
    return '\n'.join(['| '+' | '.join(heads)+' |','| '+' | '.join(['---']*len(heads))+' |']+['| '+' | '.join(map(str,r))+' |' for r in rows])
desc,tests,pairs,loso=[read(x) for x in ['anova_descriptives.csv','anova_tests.csv','anova_posthoc.csv','loso_estimates.csv']]
blocks=[('B1 Study 2 item descriptives',['Outcome','Condition','N','Mean','SD','95% CI'],
         [[r['outcome'],r['condition'],r['n'],f(r['mean']),f(r['sd']),interval(r)] for r in desc],
         'Student additional analysis. Each interval is an individual within-group t interval, not a simultaneous pairwise interval.'),
        ('B2 Study 2 omnibus tests',['Outcome','Test','F','df1','df2','p'],
         [[r['outcome'],label,f(r[prefix+'F']),f(r[prefix+'df_between'],2),f(r[prefix+'df_within'],2),pcell(r[prefix+'p'])] for r in tests for label,prefix in [('Classical',''),('Welch','welch_')]],
         'Student additional analysis. Holm correction across the two classical tests, and separately across the two Welch sensitivity tests, retains both rejections. Exact p-values, adjusted p-values, and Levene diagnostics are in anova_tests.csv.'),
        ('B3 Study 2 pairwise comparisons',['Outcome','Contrast','Difference','Family 95% CI','Holm p'],
         [[r['outcome'],r['contrast'],f(r['difference']),interval(r,'ci_low_family95','ci_high_family95'),pcell(r['p_holm_six_tests'])] for r in pairs],
         'Student additional analysis. Unequal-variance Welch comparisons. Holm p-values control the family of all six pairwise tests across both outcomes. Intervals use Bonferroni correction over the same six contrasts and can be more conservative than Holm decisions. Positive differences favor the first condition.'),
        ('B4 Study fixed effects and study omissions',['Sample','N','b','OLS SE','95% CI','p'],
         [[r['sample'],r['n'],f(r['b']),f(r['se_ols']),interval(r),pcell(r['p'])] for r in loso],
         'Student additional analysis. All models use the original five-item composite, study fixed effects, a range indicator, conventional OLS standard errors, and t-based intervals. Omission models assess study influence; they are not independent confirmatory tests.')]
for title,heads,rows,note in blocks:source+='\n\n## '+title+'\n\n'+mdtable(heads,rows)+'\n\n'+note
source+='\n\n## Figure B1 Study 2 trust and perceived accuracy\n\n![Study 2 means](additional_robustness_checks/figures/study2_trust_accuracy.png)\n\nStudent additional analysis. Points are condition means; error bars are individual 95% t intervals.\n\n## Figure B2 Leave one study out\n\n![Study omission coefficients](additional_robustness_checks/figures/leave_one_study_out.png)\n\nStudent additional analysis. Points are range coefficients and error bars are conventional OLS 95% t intervals. The dashed line marks the full-sample estimate.\n'
(ROOT/'replication_report.md').write_text(source+'\n')

doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
for name in ['Normal','Title','Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(12)
doc.styles['Normal'].paragraph_format.line_spacing=1.5
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size in [('Title',17),('Heading 1',14),('Heading 2',12)]:doc.styles[name].font.size=Pt(size)
footer=section.footer.paragraphs[0];footer.alignment=2
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def paragraph(text):
    p=doc.add_paragraph()
    for i,part in enumerate(re.split(r'\*([^*]+)\*',text)):
        run=p.add_run(part);run.italic=i%2==1
    if text.startswith(('All 31 targeted','Satisfaction is the mean','All displayed means','Student additional analysis.')):
        p.paragraph_format.line_spacing=1.1
        for run in p.runs:run.font.size=Pt(10)
    if text.startswith('Several details required attention.'):
        p.paragraph_format.page_break_before=True
    return p

def addtable(lines):
    data=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    data=[data[0]]+data[2:]
    ncols=len(data[0]);t=doc.add_table(rows=0,cols=ncols);t.autofit=False
    # Allocate space by text length so descriptive columns can wrap naturally.
    lengths=[max(len(row[i]) for row in data) for i in range(ncols)]
    weights=[min(30,max(8,x)) for x in lengths]
    widths=[6.5*x/sum(weights) for x in weights]
    if data[0][:2] == ['Study','Comparison']:
        widths=[.5,2.6,.7,.6,.5,1.6]
    for column,width in zip(t.columns,widths):column.width=Inches(width)
    for ri,values in enumerate(data):
        row=t.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for i,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[i]);cell.text=value
            tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');tcpr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

# Interpret the Markdown report, including original appendix tables.
for block in source.split('\n\n'):
    block=block.strip()
    if not block:continue
    if block.startswith('# '):doc.add_paragraph(block[2:],'Title')
    elif block.startswith('## '):
        title=block[3:]
        if not title.startswith(('1 Paper','A3 ','B2 ','B4 ')) and title!='Data and reproducibility note':doc.add_page_break()
        doc.add_paragraph(title,'Heading 1')
    elif block.startswith('### '):
        p=doc.add_paragraph(block[4:],'Heading 2')
        if block.startswith('### Pooled study'):p.paragraph_format.page_break_before=True
    elif block.startswith('|'):addtable(block.splitlines())
    elif block.startswith('!['):
        target=re.search(r'\]\((.*?)\)',block).group(1)
        doc.add_picture(str(ROOT/target),width=Inches(6.5))
    else:paragraph(block)
doc.save(ROOT/'replication_report.docx')
print('Built replication_report.docx and current additional-analysis discussion.')
